"""Read the original workbooks without changing them or requiring Excel."""

from decimal import Decimal, ROUND_HALF_UP
import hashlib
import math
from pathlib import Path
import re

from openpyxl import load_workbook
import pandas as pd

DIMENSIONS = ("traction", "roughness", "coarseness", "hardness", "temperature")
RATING_CODES = ("T", "R", "F", "H", "T*")
PARTICIPANT = re.compile(r"Participant\s*#\s*(\d+)")
QUESTION = re.compile(r"Q(\d+)_(\d+)")
FORCE_REFERENCE = re.compile(
    r"Folder:\s*([A-Za-z0-9_-]+),\s*File:\s*(force_data_(\d+)\.csv),"
    r"\s*Average Force:\s*([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)\s*"
)


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def number(value, location):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Expected a finite number at {location}; got {value!r}")
    return float(value)


def rounded_speed(value):
    """Match Excel ROUND(value, 2), including half-away-from-zero ties."""
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def parse_force_reference(value):
    match = FORCE_REFERENCE.fullmatch(str(value).strip())
    if not match:
        raise ValueError(f"Invalid Force reference: {value!r}")
    session, filename, texture, reported_force = match.groups()
    force = float(reported_force)
    if not math.isfinite(force):
        raise ValueError(f"Force reference contains a non-finite value: {value!r}")
    return session, filename, int(texture), force


def _sheet(path, *, data_only):
    book = load_workbook(path, read_only=True, data_only=data_only)
    try:
        if "Sheet1" not in book.sheetnames:
            raise ValueError(f"{path.name}: Sheet1 is missing")
        return list(book["Sheet1"].values)
    finally:
        book.close()


def _participants(header):
    result = {}
    for column, value in enumerate(header):
        match = PARTICIPANT.fullmatch(str(value or "").strip())
        if match:
            participant = int(match.group(1))
            if participant in result:
                raise ValueError(f"Duplicate participant column: {participant}")
            result[participant] = column
    if sorted(result) != list(range(1, 41)):
        raise ValueError("Expected exactly Participant #1 through #40")
    return result


def load_observations(data_dir):
    """Return 400 participant/texture observations and a workbook audit.

    NF is extracted from archived Force text, and SS is recomputed from Speed.
    F retains the questionnaire's -10=fine, +10=coarse polarity.
    Texture/force numbering follows workbook alignment, not a plate registry.
    """
    data_dir = Path(data_dir)
    combined_path = data_dir / "processed" / "All in one.xlsx"
    ratings_path = data_dir / "processed" / "Perceptual Dimension Ratings.xlsx"
    for path in (combined_path, ratings_path):
        if not path.is_file():
            raise FileNotFoundError(f"Missing input: {path}. Run from the repository root or set --data-dir.")
    combined = _sheet(combined_path, data_only=True)
    formulas = _sheet(combined_path, data_only=False)
    ratings = _sheet(ratings_path, data_only=True)
    cp, rp = _participants(combined[0]), _participants(ratings[0])
    if len(combined) < 91 or len(ratings) < 51:
        raise ValueError("Workbook has fewer rows than the documented source layout")
    for texture in range(1, 11):
        for offset, code in enumerate(RATING_CODES):
            if combined[1 + (texture - 1) * 5 + offset][0] != code:
                raise ValueError(f"Unexpected rating row labels for texture {texture}")
        for prefix, start in (("NF", 51), ("SS", 61), ("Force", 71), ("Speed", 81)):
            if combined[start + texture - 1][0] != f"{prefix}{texture}":
                raise ValueError(f"Missing row label {prefix}{texture}")

    questions, materials = {}, {}
    for row in ratings[1:51]:
        match = QUESTION.fullmatch(str(row[0] or ""))
        if not match:
            raise ValueError(f"Unknown questionnaire item: {row[0]!r}")
        texture, dimension = map(int, match.groups())
        if not 1 <= texture <= 10 or not 1 <= dimension <= 5 or (texture, dimension) in questions:
            raise ValueError(f"Duplicate or out-of-range questionnaire item: {row[0]}")
        questions[texture, dimension] = row
        if row[1] is not None:
            material = str(row[1]).strip()
            if texture in materials and materials[texture] != material:
                raise ValueError(f"Inconsistent material name for texture {texture}")
            materials[texture] = material
    if len(questions) != 50 or len(materials) != 10:
        raise ValueError("Expected 10 materials with 5 dimensions each")
    if "-10: very fine, +10: very coarse" not in str(questions[1, 3][2]):
        raise ValueError("Questionnaire coarseness polarity differs from the documented source")

    observations = []
    cached_matches = {"NF": 0, "SS": 0}
    missing_caches = []
    for participant, column in sorted(cp.items()):
        sessions = set()
        for texture in range(1, 11):
            session, filename, force_texture, force = parse_force_reference(combined[70 + texture][column])
            if force_texture != texture:
                raise ValueError(f"Force filename disagrees with texture row: {filename}")
            sessions.add(session)
            raw_speed = number(combined[80 + texture][column], f"Speed{texture}/P{participant}")
            speed = rounded_speed(raw_speed)
            row = {"participant_id": f"P{participant:02}", "session_id": session,
                   "texture_id": texture, "texture": materials[texture]}
            for dimension, name in enumerate(DIMENSIONS, 1):
                value = number(questions[texture, dimension][rp[participant]], f"Q{texture}_{dimension}/P{participant}")
                if not -10 <= value <= 10:
                    raise ValueError("Questionnaire rating is outside [-10, 10]")
                archived = number(combined[(texture - 1) * 5 + dimension][column], "combined rating")
                if value != archived:
                    raise ValueError(f"Workbooks disagree at P{participant}/texture {texture}/{name}")
                row[name] = value
            for prefix, index, expected in (("NF", 50 + texture, force), ("SS", 60 + texture, speed)):
                cached = combined[index][column]
                if cached is None and str(formulas[index][column]).startswith("="):
                    missing_caches.append(f"{prefix}{texture}/P{participant}")
                elif abs(number(cached, f"{prefix}{texture}/P{participant}") - expected) > 1e-8:
                    raise ValueError(f"{prefix} cached value disagrees with its source at P{participant}/texture {texture}")
                else:
                    cached_matches[prefix] += 1
            row.update(reported_force=force, reported_speed=speed, speed_unrounded=raw_speed,
                       source_force_csv=f"raw/force/{session}/{filename}")
            observations.append(row)
        if len(sessions) != 1:
            raise ValueError(f"Participant {participant} maps to multiple raw sessions")
    frame = pd.DataFrame(observations)
    if frame["session_id"].nunique() != 40 or frame["source_force_csv"].nunique() != 400:
        raise ValueError("Expected a unique raw session per participant and 400 unique trial files")
    audit = {
        "participants": 40, "textures": 10, "observations": len(frame), "ratings_compared": 2000,
        "ratings_identical": 2000, "cached_formula_matches": cached_matches,
        "missing_formula_caches_reconstructed": missing_caches,
        "texture_alignment": "Force/Speed row numbers aligned with questionnaire order; no independent plate registry",
        "coarseness": "Source question calls this Fineness; -10=fine, +10=coarse; values are not inverted",
        "units": {
            "reported_force": "Not specified in the archived processed workbook",
            "reported_speed": "Author confirms pixel-to-mm spatial calibration; factor/file unavailable; time unit unconfirmed",
        },
        "input_sha256": {f"processed/{p.name}": sha256(p) for p in (combined_path, ratings_path)},
    }
    return frame, audit
